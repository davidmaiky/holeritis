"""
Gerador de Planilhas Excel (.xlsx) Profissionais e Detalhadas para Holerites e Folha de Pagamento.
Inclui:
- Todos os dados cadastrais da empresa e de cada colaborador
- Demonstração completa de todos os eventos (Créditos em Azul e Débitos em Vermelho)
- Resumo financeiro consolidado e quadro de bases de cálculo legais
- Duas abas: 'Holerites Detalhados' e 'Resumo Consolidado'
"""

import io
import re
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def gerar_workbook_excel(rel, itens_filtrados=None):
    """
    Gera um objeto openpyxl.Workbook formatado com o relatório completo da folha.
    itens_filtrados: lista opcional de itens a exportar (caso venha filtrado da interface).
    """
    wb = openpyxl.Workbook()

    # Tipografia e Fontes
    font_title = Font(name="Segoe UI", size=13, bold=True, color="FFFFFF")
    font_subtitle = Font(name="Segoe UI", size=10, italic=True, color="FFFFFF")
    font_meta = Font(name="Segoe UI", size=9, color="475569")
    font_th = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    
    font_emp_header = Font(name="Segoe UI", size=11, bold=True, color="0F172A")
    font_emp_sub = Font(name="Segoe UI", size=9, bold=False, color="334155")
    
    font_text = Font(name="Segoe UI", size=9, color="1F2937")
    font_code = Font(name="Consolas", size=9, color="475569")
    font_muted = Font(name="Segoe UI", size=9, color="9CA3AF")
    
    # Cores de Crédito (Azul Real / Royal Blue)
    font_credito_badge = Font(name="Segoe UI", size=9, bold=True, color="1E40AF")
    font_credito_val = Font(name="Segoe UI", size=9, bold=True, color="1E40AF")
    fill_credito_badge = PatternFill(start_color="DBEAFE", end_color="DBEAFE", fill_type="solid")
    fill_credito_cell = PatternFill(start_color="EFF6FF", end_color="EFF6FF", fill_type="solid")
    border_credito = Border(
        left=Side(style="thin", color="BFDBFE"),
        right=Side(style="thin", color="BFDBFE"),
        top=Side(style="thin", color="BFDBFE"),
        bottom=Side(style="thin", color="BFDBFE")
    )

    # Cores de Débito (Vermelho Carmesim / Crimson)
    font_debito_badge = Font(name="Segoe UI", size=9, bold=True, color="991B1B")
    font_debito_val = Font(name="Segoe UI", size=9, bold=True, color="B91C1C")
    fill_debito_badge = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
    fill_debito_cell = PatternFill(start_color="FEF2F2", end_color="FEF2F2", fill_type="solid")
    border_debito = Border(
        left=Side(style="thin", color="FECACA"),
        right=Side(style="thin", color="FECACA"),
        top=Side(style="thin", color="FECACA"),
        bottom=Side(style="thin", color="FECACA")
    )

    # Valor Líquido (Verde Institucional)
    font_liquido_val = Font(name="Segoe UI", size=11, bold=True, color="166534")
    fill_liquido = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")

    font_totais_header = Font(name="Segoe UI", size=10, bold=True, color="0F172A")
    fill_totais_bar = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")

    fill_brand = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    fill_table_head = PatternFill(start_color="334155", end_color="334155", fill_type="solid")
    fill_emp_head = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
    fill_bases_head = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

    thin_gray = Side(style="thin", color="CBD5E1")
    border_box = Border(left=thin_gray, right=thin_gray, top=thin_gray, bottom=thin_gray)

    num_fmt_curr = '"R$" #,##0.00;[Red]-"R$" #,##0.00;"R$" 0.00'

    itens = itens_filtrados if itens_filtrados is not None else rel.get("itens", [])

    # =========================================================================
    # ABA 1: HOLERITES DETALHADOS (Relatório Completo de Cada Funcionário)
    # =========================================================================
    ws1 = wb.active
    ws1.title = "Holerites Detalhados"
    ws1.views.sheetView[0].showGridLines = True

    # Banner Institucional da Empresa
    ws1.merge_cells("A1:F1")
    empresa_nome = rel.get("empresa") or "EMPRESA"
    ws1["A1"] = f"DEMONSTRATIVO DETALHADO DE FOLHA DE PAGAMENTO - {empresa_nome}"
    ws1["A1"].font = font_title
    ws1["A1"].fill = fill_brand
    ws1["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws1.row_dimensions[1].height = 28

    ws1.merge_cells("A2:F2")
    cnpj_txt = f"CNPJ: {rel.get('cnpj', 'Não informado')}" if rel.get("cnpj") else ""
    periodo_txt = f"Período de Apuração: {rel.get('periodo_texto', rel.get('mes_ano', ''))}"
    ws1["A2"] = f"{cnpj_txt}   |   {periodo_txt}   |   Competência: {rel.get('mes_ano', '')}"
    ws1["A2"].font = font_subtitle
    ws1["A2"].fill = fill_brand
    ws1["A2"].alignment = Alignment(horizontal="center", vertical="center")
    ws1.row_dimensions[2].height = 20

    current_row = 4

    for emp in itens:
        dados = emp.get("dados_adicionais") or {}
        bases = emp.get("bases") or {}
        eventos = emp.get("eventos") or []

        # Card do Funcionário - Linha 1: Nome e Código
        ws1.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=6)
        cell_emp = ws1.cell(row=current_row, column=1)
        cod_str = f"Cód: {emp.get('codigo', '-')}" if emp.get('codigo') else "Cód: -"
        cell_emp.value = f"👤 {emp.get('nome', '')}  ({cod_str})"
        cell_emp.font = font_emp_header
        cell_emp.fill = fill_emp_head
        cell_emp.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        for col_idx in range(1, 7):
            ws1.cell(row=current_row, column=col_idx).border = border_box
        ws1.row_dimensions[current_row].height = 24
        current_row += 1

        # Card do Funcionário - Linha 2: Metadados
        ws1.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=6)
        cell_meta = ws1.cell(row=current_row, column=1)
        sal_base_fmt = f"R$ {emp.get('salario', 0):,.2f}".replace(',', 'X').replace('.', ',').replace('X', '.')
        meta_info = (
            f"Cargo: {emp.get('funcao') or 'Não informado'}   |   "
            f"Admissão: {dados.get('admissao') or '-'}   |   "
            f"Situação: {dados.get('situacao') or 'Ativo'}   |   "
            f"Dep. IR: {dados.get('dependentes_ir', 0)}   |   "
            f"Salário Base: {sal_base_fmt}"
        )
        cell_meta.value = meta_info
        cell_meta.font = font_emp_sub
        cell_meta.fill = fill_bases_head
        cell_meta.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        for col_idx in range(1, 7):
            ws1.cell(row=current_row, column=col_idx).border = border_box
        ws1.row_dimensions[current_row].height = 20
        current_row += 1

        # Cabeçalho da Tabela de Rubricas
        headers_events = ["Cód.", "Descrição da Rubrica / Evento", "Referência", "Tipo", "Crédito / Proventos (R$)", "Débito / Descontos (R$)"]
        for c_idx, h_text in enumerate(headers_events, start=1):
            c = ws1.cell(row=current_row, column=c_idx, value=h_text)
            c.font = font_th
            c.fill = fill_table_head
            c.alignment = Alignment(horizontal="center" if c_idx in [1, 3, 4] else ("right" if c_idx in [5, 6] else "left"), vertical="center")
            c.border = border_box
        ws1.row_dimensions[current_row].height = 22
        current_row += 1

        # Extrair e ordenar rubricas
        evts_to_render = eventos if eventos and len(eventos) > 0 else [
            {"codigo": "001", "descricao": "Salário Base / Proventos Contratuais", "referencia": "30d", "valor": emp.get("proventos", 0), "tipo": "provento"}
        ]
        if not eventos and emp.get("adiantamento_anterior", 0) > 0:
            evts_to_render.append({"codigo": "012", "descricao": "Adiantamento Anterior Compensado", "referencia": "-", "valor": emp.get("adiantamento_anterior", 0), "tipo": "desconto"})
        if not eventos and emp.get("descontos", 0) > 0:
            evts_to_render.append({"codigo": "999", "descricao": "Descontos e Retenções Legais Totais", "referencia": "-", "valor": emp.get("descontos", 0), "tipo": "desconto"})

        soma_creditos = 0.0
        soma_debitos = 0.0

        for evt in evts_to_render:
            is_prov = (evt.get("tipo") == "provento")
            val = float(evt.get("valor") or 0.0)

            c_cod = ws1.cell(row=current_row, column=1, value=str(evt.get("codigo") or "-"))
            c_cod.font = font_code
            c_cod.alignment = Alignment(horizontal="center", vertical="center")
            c_cod.border = border_box

            c_desc = ws1.cell(row=current_row, column=2, value=str(evt.get("descricao") or "Item"))
            c_desc.font = font_text
            c_desc.alignment = Alignment(horizontal="left", vertical="center")
            c_desc.border = border_box

            ref_val = str(evt.get("referencia") or "-")
            c_ref = ws1.cell(row=current_row, column=3, value=ref_val if ref_val else "-")
            c_ref.font = font_code
            c_ref.alignment = Alignment(horizontal="center", vertical="center")
            c_ref.border = border_box

            c_tipo = ws1.cell(row=current_row, column=4)
            c_prov = ws1.cell(row=current_row, column=5)
            c_desc_val = ws1.cell(row=current_row, column=6)

            if is_prov:
                soma_creditos += val
                c_tipo.value = "CRÉDITO"
                c_tipo.font = font_credito_badge
                c_tipo.fill = fill_credito_badge
                c_tipo.alignment = Alignment(horizontal="center", vertical="center")
                c_tipo.border = border_credito

                c_prov.value = val
                c_prov.number_format = num_fmt_curr
                c_prov.font = font_credito_val
                c_prov.fill = fill_credito_cell
                c_prov.alignment = Alignment(horizontal="right", vertical="center")
                c_prov.border = border_credito

                c_desc_val.value = "-"
                c_desc_val.font = font_muted
                c_desc_val.alignment = Alignment(horizontal="center", vertical="center")
                c_desc_val.border = border_box
            else:
                soma_debitos += val
                c_tipo.value = "DÉBITO"
                c_tipo.font = font_debito_badge
                c_tipo.fill = fill_debito_badge
                c_tipo.alignment = Alignment(horizontal="center", vertical="center")
                c_tipo.border = border_debito

                c_prov.value = "-"
                c_prov.font = font_muted
                c_prov.alignment = Alignment(horizontal="center", vertical="center")
                c_prov.border = border_box

                c_desc_val.value = val
                c_desc_val.number_format = num_fmt_curr
                c_desc_val.font = font_debito_val
                c_desc_val.fill = fill_debito_cell
                c_desc_val.alignment = Alignment(horizontal="right", vertical="center")
                c_desc_val.border = border_debito

            ws1.row_dimensions[current_row].height = 20
            current_row += 1

        # Linha de Totais dos Eventos (Proventos x Descontos)
        ws1.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=4)
        c_lbl_tot = ws1.cell(row=current_row, column=1, value="TOTAIS DE EVENTOS (CRÉDITOS & DÉBITOS):")
        c_lbl_tot.font = font_totais_header
        c_lbl_tot.fill = fill_totais_bar
        c_lbl_tot.alignment = Alignment(horizontal="right", vertical="center")
        for col_idx in range(1, 5):
            ws1.cell(row=current_row, column=col_idx).border = border_box

        c_tot_prov = ws1.cell(row=current_row, column=5, value=soma_creditos or emp.get("proventos", 0.0))
        c_tot_prov.number_format = num_fmt_curr
        c_tot_prov.font = font_credito_val
        c_tot_prov.fill = fill_credito_badge
        c_tot_prov.alignment = Alignment(horizontal="right", vertical="center")
        c_tot_prov.border = border_credito

        c_tot_desc = ws1.cell(row=current_row, column=6, value=soma_debitos or emp.get("descontos", 0.0))
        c_tot_desc.number_format = num_fmt_curr
        c_tot_desc.font = font_debito_val
        c_tot_desc.fill = fill_debito_badge
        c_tot_desc.alignment = Alignment(horizontal="right", vertical="center")
        c_tot_desc.border = border_debito
        ws1.row_dimensions[current_row].height = 22
        current_row += 1

        # Linha de Resumo Financeiro & Líquido a Receber
        ws1.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=3)
        c_adiant_lbl = ws1.cell(row=current_row, column=1)
        ad_ant = emp.get("adiantamento_anterior", 0.0)
        c_adiant_lbl.value = f"Adiantamento Anterior Compensado: R$ {ad_ant:,.2f}".replace(',', 'X').replace('.', ',').replace('X', '.')
        c_adiant_lbl.font = font_meta
        c_adiant_lbl.fill = fill_bases_head
        c_adiant_lbl.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        for col_idx in range(1, 4):
            ws1.cell(row=current_row, column=col_idx).border = border_box

        c_liq_lbl = ws1.cell(row=current_row, column=4, value="VALOR LÍQUIDO:")
        c_liq_lbl.font = font_liquido_val
        c_liq_lbl.fill = fill_liquido
        c_liq_lbl.alignment = Alignment(horizontal="right", vertical="center")
        c_liq_lbl.border = border_box

        ws1.merge_cells(start_row=current_row, start_column=5, end_row=current_row, end_column=6)
        c_liq_val = ws1.cell(row=current_row, column=5, value=emp.get("liquido", 0.0))
        c_liq_val.number_format = num_fmt_curr
        c_liq_val.font = font_liquido_val
        c_liq_val.fill = fill_liquido
        c_liq_val.alignment = Alignment(horizontal="center", vertical="center")
        for col_idx in [5, 6]:
            ws1.cell(row=current_row, column=col_idx).border = border_box
        ws1.row_dimensions[current_row].height = 24
        current_row += 1

        # Quadro de Bases de Cálculo e Encargos
        bases_headers = [
            ("Salário Base", emp.get("salario", 0.0)),
            ("Base INSS Empresa", bases.get("base_inss_empresa", 0.0)),
            ("Base INSS Func.", bases.get("base_inss_funcionario", 0.0)),
            ("Base FGTS", bases.get("base_fgts", 0.0)),
            ("F.G.T.S. Mês (8%)", bases.get("valor_fgts", 0.0)),
            ("Base IRRF", bases.get("base_irrf", 0.0))
        ]
        for b_idx, (b_title, _) in enumerate(bases_headers, start=1):
            c_b_head = ws1.cell(row=current_row, column=b_idx, value=b_title)
            c_b_head.font = Font(name="Segoe UI", size=8, bold=True, color="64748B")
            c_b_head.fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
            c_b_head.alignment = Alignment(horizontal="center", vertical="center")
            c_b_head.border = border_box
        ws1.row_dimensions[current_row].height = 18
        current_row += 1

        for b_idx, (_, b_val) in enumerate(bases_headers, start=1):
            c_b_val = ws1.cell(row=current_row, column=b_idx, value=float(b_val or 0.0))
            c_b_val.number_format = num_fmt_curr
            c_b_val.font = Font(name="Consolas", size=8, bold=True, color="334155")
            c_b_val.fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
            c_b_val.alignment = Alignment(horizontal="center", vertical="center")
            c_b_val.border = border_box
        ws1.row_dimensions[current_row].height = 18
        current_row += 2  # Espaço para o próximo funcionário

    col_widths_1 = [10, 42, 14, 14, 22, 22]
    for i, w in enumerate(col_widths_1, start=1):
        ws1.column_dimensions[get_column_letter(i)].width = w

    # =========================================================================
    # ABA 2: RESUMO GERAL DA FOLHA (Consolidado de Conferência)
    # =========================================================================
    ws2 = wb.create_sheet(title="Resumo Consolidado")
    ws2.views.sheetView[0].showGridLines = True

    # Cabeçalho Geral
    ws2.merge_cells("A1:K1")
    ws2["A1"] = f"RESUMO CONSOLIDADO DA FOLHA - {empresa_nome}"
    ws2["A1"].font = font_title
    ws2["A1"].fill = fill_brand
    ws2["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws2.row_dimensions[1].height = 28

    ws2.merge_cells("A2:K2")
    ws2["A2"] = f"Competência: {rel.get('mes_ano', '')}   |   {periodo_txt}   |   Total de Colaboradores: {len(itens)}"
    ws2["A2"].font = font_subtitle
    ws2["A2"].fill = fill_brand
    ws2["A2"].alignment = Alignment(horizontal="center", vertical="center")
    ws2.row_dimensions[2].height = 20

    headers_summary = [
        "Cód.", "Colaborador", "Cargo / Função", "Salário Base (R$)",
        "Créditos / Proventos (R$)", "Adiantamento Anterior (R$)", "Débitos / Descontos (R$)",
        "Total Líquido (R$)", "Base INSS (R$)", "Base FGTS (R$)", "FGTS do Mês (R$)"
    ]
    for c_idx, h_text in enumerate(headers_summary, start=1):
        c = ws2.cell(row=4, column=c_idx, value=h_text)
        c.font = font_th
        c.fill = fill_table_head
        c.alignment = Alignment(horizontal="center" if c_idx in [1] else ("right" if c_idx >= 4 else "left"), vertical="center")
        c.border = border_box
    ws2.row_dimensions[4].height = 24

    r_idx = 5
    for emp in itens:
        bases = emp.get("bases") or {}
        ws2.cell(row=r_idx, column=1, value=str(emp.get("codigo") or "-")).alignment = Alignment(horizontal="center", vertical="center")
        ws2.cell(row=r_idx, column=2, value=str(emp.get("nome") or "")).alignment = Alignment(horizontal="left", vertical="center")
        ws2.cell(row=r_idx, column=3, value=str(emp.get("funcao") or "")).alignment = Alignment(horizontal="left", vertical="center")

        c_sb = ws2.cell(row=r_idx, column=4, value=float(emp.get("salario") or 0.0))
        c_sb.number_format = num_fmt_curr
        c_sb.alignment = Alignment(horizontal="right", vertical="center")

        # Proventos em AZUL
        c_pr = ws2.cell(row=r_idx, column=5, value=float(emp.get("proventos") or 0.0))
        c_pr.number_format = num_fmt_curr
        c_pr.font = font_credito_val
        c_pr.fill = fill_credito_cell
        c_pr.alignment = Alignment(horizontal="right", vertical="center")
        c_pr.border = border_credito

        c_ad = ws2.cell(row=r_idx, column=6, value=float(emp.get("adiantamento_anterior") or 0.0))
        c_ad.number_format = num_fmt_curr
        c_ad.alignment = Alignment(horizontal="right", vertical="center")

        # Descontos em VERMELHO
        c_ds = ws2.cell(row=r_idx, column=7, value=float(emp.get("descontos") or 0.0))
        c_ds.number_format = num_fmt_curr
        c_ds.font = font_debito_val
        c_ds.fill = fill_debito_cell
        c_ds.alignment = Alignment(horizontal="right", vertical="center")
        c_ds.border = border_debito

        # Líquido em VERDE
        c_lq = ws2.cell(row=r_idx, column=8, value=float(emp.get("liquido") or 0.0))
        c_lq.number_format = num_fmt_curr
        c_lq.font = font_liquido_val
        c_lq.fill = fill_liquido
        c_lq.alignment = Alignment(horizontal="right", vertical="center")

        c_binss = ws2.cell(row=r_idx, column=9, value=float(bases.get("base_inss_funcionario") or 0.0))
        c_binss.number_format = num_fmt_curr
        c_binss.alignment = Alignment(horizontal="right", vertical="center")

        c_bfgts = ws2.cell(row=r_idx, column=10, value=float(bases.get("base_fgts") or 0.0))
        c_bfgts.number_format = num_fmt_curr
        c_bfgts.alignment = Alignment(horizontal="right", vertical="center")

        c_vfgts = ws2.cell(row=r_idx, column=11, value=float(bases.get("valor_fgts") or 0.0))
        c_vfgts.number_format = num_fmt_curr
        c_vfgts.alignment = Alignment(horizontal="right", vertical="center")

        for c_c in range(1, 12):
            cell = ws2.cell(row=r_idx, column=c_c)
            if cell.font == font_text or cell.font.name != "Segoe UI":
                cell.font = font_text
            if not cell.border.top.style:
                cell.border = border_box

        ws2.row_dimensions[r_idx].height = 20
        r_idx += 1

    # Linha de TOTAIS CONSOLIDADOS
    ws2.merge_cells(start_row=r_idx, start_column=1, end_row=r_idx, end_column=3)
    c_tot_label = ws2.cell(row=r_idx, column=1, value=f"TOTAIS CONSOLIDADOS ({len(itens)} colaboradores)")
    c_tot_label.font = Font(name="Segoe UI", size=10, bold=True, color="0F172A")
    c_tot_label.fill = fill_totais_bar
    c_tot_label.alignment = Alignment(horizontal="right", vertical="center")
    for col_c in range(1, 4):
        ws2.cell(row=r_idx, column=col_c).border = border_box

    tot_cols = [
        (4, sum(float(x.get("salario") or 0) for x in itens), None, None),
        (5, sum(float(x.get("proventos") or 0) for x in itens), font_credito_val, fill_credito_badge),
        (6, sum(float(x.get("adiantamento_anterior") or 0) for x in itens), None, None),
        (7, sum(float(x.get("descontos") or 0) for x in itens), font_debito_val, fill_debito_badge),
        (8, sum(float(x.get("liquido") or 0) for x in itens), font_liquido_val, fill_liquido),
        (9, sum(float((x.get("bases") or {}).get("base_inss_funcionario") or 0) for x in itens), None, None),
        (10, sum(float((x.get("bases") or {}).get("base_fgts") or 0) for x in itens), None, None),
        (11, sum(float((x.get("bases") or {}).get("valor_fgts") or 0) for x in itens), None, None),
    ]

    for col_num, val, f_custom, fill_custom in tot_cols:
        cell_tot = ws2.cell(row=r_idx, column=col_num, value=val)
        cell_tot.number_format = num_fmt_curr
        cell_tot.font = f_custom or Font(name="Segoe UI", size=10, bold=True, color="0F172A")
        cell_tot.fill = fill_custom or fill_totais_bar
        cell_tot.alignment = Alignment(horizontal="right", vertical="center")
        cell_tot.border = border_box

    ws2.row_dimensions[r_idx].height = 24

    col_widths_2 = [8, 36, 30, 18, 22, 22, 22, 22, 18, 18, 18]
    for i, w in enumerate(col_widths_2, start=1):
        ws2.column_dimensions[get_column_letter(i)].width = w

    return wb

def gerar_excel_bytes(rel, itens_filtrados=None):
    """
    Retorna os bytes do arquivo Excel (.xlsx).
    """
    wb = gerar_workbook_excel(rel, itens_filtrados)
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()
